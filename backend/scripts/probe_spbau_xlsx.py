"""Скачать и сохранить структуру xlsx списка СПбАУ в файл."""
import io
import json
import urllib.request

import openpyxl

url = "https://www.spbau.ru/images/documents/campaign/bachelor/lists/B0807.xlsx"
data = urllib.request.urlopen(url).read()
wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)

out = {"sheets": wb.sheetnames, "rows": {}}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([str(c) if c is not None else None for c in row])
    out["rows"][sheet_name] = rows

with open("scripts/spbau_xlsx_sample.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)

print("saved", len(out["rows"][wb.sheetnames[0]]), "rows")
