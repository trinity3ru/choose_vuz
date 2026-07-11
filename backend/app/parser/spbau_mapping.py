"""
Преобразование xlsx-списка Академического университета им. Алферова (spbau.ru).

Сайт публикует конкурсные списки одним xlsx-файлом на странице
/campaign/bachelor/documents. Файл содержит несколько конкурсных групп;
для направления 03.03.01 объединяем все программы с «Общий конкурс» + «Очная» + _бюджет.
"""

import io
import re
from urllib.parse import urljoin

import openpyxl
from bs4 import BeautifulSoup

from app.schemas.parser_schema import ApplicantRow

# Колонки таблицы абитуриентов (0-based).
_COL_RANK = 0
_COL_EPGU = 2
_COL_PRIORITY = 4
_COL_EXAM_TOTAL = 10
_COL_ACHIEVEMENT = 11
_COL_TARGET_ACHIEVEMENT = 12
_COL_TOTAL = 13
_COL_ACHIEVEMENTS_TEXT = 14
_COL_PREF_9 = 17
_COL_PREF_10 = 18
_COL_AGREEMENT = 19
_COL_E_AGREEMENT = 20
_COL_NOTE = 24

# Минимальное число колонок в строке данных.
_MIN_COLS = 21

# Маркеры подразделов таблицы — не являются строками абитуриентов.
_SECTION_MARKERS = {
    "участвуют в конкурсе",
    "по результатам экзаменов",
    "без экзаменов",
    "не участвуют в конкурсе",
    "1 уровень олимпиады, призер",
    "1 уровень олимпиады, победитель",
}

_RANK_RE = re.compile(r"^\d+$")
_EPGU_RE = re.compile(r"^\d{5,}$")


def to_int(value: object) -> int | None:
    """Привести значение ячейки к целому (пустое -> None)."""
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except (ValueError, TypeError):
        return None


def _cell(row: tuple, index: int) -> str:
    """Безопасно получить текст ячейки строки."""
    if index >= len(row):
        return ""
    value = row[index]
    return "" if value is None else str(value).strip()


def _is_yes(value: str) -> bool:
    return value.strip().lower() in {"да", "+"}


def find_latest_xlsx_url(html: str, page_url: str) -> str:
    """
    Найти ссылку на актуальный xlsx «Основной конкурс» на странице документов.

    Берём первую ссылку на .xlsx в таблице (самая свежая дата сверху).
    """
    soup = BeautifulSoup(html, "html.parser")
    for link in soup.select('a[href$=".xlsx"]'):
        href = link.get("href", "")
        if "lists" in href.lower():
            return urljoin(page_url, href)

    # Запасной вариант: любая xlsx-ссылка на странице.
    for link in soup.select('a[href$=".xlsx"]'):
        href = link.get("href", "")
        if href:
            return urljoin(page_url, href)

    raise RuntimeError("на странице документов не найдена ссылка на xlsx")


def _is_data_row(row: tuple) -> bool:
    """Строка с данными абитуриента: № п/п — число, УИД — числовой код ЕПГУ."""
    rank = _cell(row, _COL_RANK)
    if not _RANK_RE.match(rank):
        return False
    if rank.lower() in _SECTION_MARKERS:
        return False

    epgu = _cell(row, _COL_EPGU)
    if not _EPGU_RE.match(epgu):
        return False

    return len(row) >= _MIN_COLS


def _section_matches(study_form: str, finance_type: str, meta: dict[str, str | int | None]) -> bool:
    """
    Проверить, подходит ли текущий блок xlsx под параметры направления из конфига.

    На сайте: «Общий конкурс» + «Очная» + суффикс «_бюджет» в названии группы.
    """
    group = str(meta.get("group") or "")
    form = str(meta.get("form") or "")
    category = str(meta.get("category") or "")

    if form != study_form:
        return False

    if finance_type == "Бюджетная основа":
        return category == "Общий конкурс" and group.endswith("_бюджет")

    return False


def row_to_applicant(row: tuple) -> ApplicantRow:
    """Собрать ApplicantRow из строки xlsx."""
    achievements = _cell(row, _COL_ACHIEVEMENTS_TEXT)
    note = _cell(row, _COL_NOTE)
    pref9 = _cell(row, _COL_PREF_9)
    pref10 = _cell(row, _COL_PREF_10)
    preferential = "Да" if (_is_yes(pref9) or _is_yes(pref10)) else None

    combined = f"{achievements} {note}".lower()
    is_bvi = "олимпиад" in combined or "бви" in combined or "без экзамен" in combined

    has_agreement = _is_yes(_cell(row, _COL_AGREEMENT)) or _is_yes(_cell(row, _COL_E_AGREEMENT))

    return ApplicantRow(
        applicant_code=_cell(row, _COL_EPGU),
        is_bvi=is_bvi,
        total_score=to_int(row[_COL_TOTAL] if len(row) > _COL_TOTAL else None),
        exam_score=to_int(row[_COL_EXAM_TOTAL] if len(row) > _COL_EXAM_TOTAL else None),
        achievement_score=to_int(row[_COL_ACHIEVEMENT] if len(row) > _COL_ACHIEVEMENT else None),
        target_achievement_score=to_int(
            row[_COL_TARGET_ACHIEVEMENT] if len(row) > _COL_TARGET_ACHIEVEMENT else None
        ),
        preferential_right=preferential,
        priority=to_int(row[_COL_PRIORITY] if len(row) > _COL_PRIORITY else None),
        has_agreement=has_agreement,
        # В БД review_status — varchar(200). Колонка «Достижения» бывает длинной,
        # поэтому сохраняем только краткое «Примечание» (как статус/комментарий).
        review_status=note or None,
    )


def parse_xlsx(
    content: bytes,
    study_form: str,
    finance_type: str,
) -> tuple[int | None, list[ApplicantRow], str | None]:
    """
    Разобрать xlsx и вернуть объединённый список по фильтру формы/основы.

    :return: (сумма мест по блокам, абитуриенты, дата/время из имени файла — опционально)
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    meta: dict[str, str | int | None] = {}
    applicants: list[ApplicantRow] = []
    places_total = 0
    section_active = False

    for row in ws.iter_rows(values_only=True):
        label = _cell(row, 1)

        if label == "Конкурсная группа:":
            meta["group"] = _cell(row, 2)
            section_active = False
            continue
        if label == "Форма обучения:":
            meta["form"] = _cell(row, 2)
            section_active = False
            continue
        if label == "Категория:":
            meta["category"] = _cell(row, 2)
            section_active = False
            continue
        if label == "План приема:":
            meta["places"] = to_int(row[2] if len(row) > 2 else None)
            section_active = _section_matches(study_form, finance_type, meta)
            if section_active and meta.get("places"):
                places_total += int(meta["places"])
            continue

        if not section_active:
            continue

        first = _cell(row, _COL_RANK).lower()
        if first in _SECTION_MARKERS:
            continue

        if _is_data_row(row):
            applicants.append(row_to_applicant(row))

    wb.close()

    places = places_total if places_total > 0 else None
    return places, applicants, None
